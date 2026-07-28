#!/usr/bin/env python3
"""
Merge a batch's input (master data) with the enrichment result JSON files and
write batches/Batch_XXX.xlsx.

Result files: state/results/row_<row_id>.json  (one per doctor, written by agents)
Each must conform to the RESULT_KEY_TO_COLUMN contract in schema.py.

Usage: python3 scripts/write_batch_xlsx.py <batch_no>
"""
import json, os, sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(__file__))
from schema import OUTPUT_COLUMNS, RESULT_KEY_TO_COLUMN

def load_result(row_id):
    p = f"state/results/row_{row_id}.json"
    if not os.path.exists(p):
        return None
    with open(p) as f:
        return json.load(f)

def build_row(rec, res):
    row = {c: "" for c in OUTPUT_COLUMNS}
    row["Doctor Name"] = rec["name"]
    row["Department"] = rec["department"]
    row["Experience"] = rec["experience_raw"]
    row["Hospital"] = rec["hospital"]
    row["Address"] = rec["address"]
    row["City"] = rec.get("city_guess", "")
    row["Priority Rank"] = rec["priority_rank"]
    row["HNI Priority Score"] = rec["priority_score"]
    if res:
        for k, col in RESULT_KEY_TO_COLUMN.items():
            v = res.get(k, "")
            if isinstance(v, bool):
                v = "Yes" if v else ""
            row[col] = v if v is not None else ""
        # prefer verified city if provided
        if res.get("city"):
            row["City"] = res["city"]
    else:
        row["Status"] = "Not yet processed"
    return row

def main():
    batch_no = int(sys.argv[1])
    with open(f"state/batch_{batch_no:03d}_input.json") as f:
        slice_ = json.load(f)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Batch_{batch_no:03d}"

    # header
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for j, col in enumerate(OUTPUT_COLUMNS, 1):
        c = ws.cell(row=1, column=j, value=col)
        c.fill = hdr_fill; c.font = hdr_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    processed = 0
    for i, rec in enumerate(slice_, start=2):
        res = load_result(rec["row_id"])
        if res: processed += 1
        row = build_row(rec, res)
        for j, col in enumerate(OUTPUT_COLUMNS, 1):
            c = ws.cell(row=i, column=j, value=row[col])
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = border
            c.font = Font(size=9)

    # column widths
    widths = {"Doctor Name":24,"Department":22,"Experience":11,"Hospital":18,
              "Address":30,"City":16,"Priority Rank":8,"HNI Priority Score":9,
              "Sources":50,"Notes":40,"HNI Signals":34,"Verified Role/Title":26,
              "Status":18}
    for j, col in enumerate(OUTPUT_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(col, 24)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 42

    os.makedirs("batches", exist_ok=True)
    out = f"batches/Batch_{batch_no:03d}.xlsx"
    wb.save(out)
    print(f"Wrote {out}: {len(slice_)} doctors, {processed} enriched, "
          f"{len(slice_)-processed} pending")

if __name__ == "__main__":
    main()
